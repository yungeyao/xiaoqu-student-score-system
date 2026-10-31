"""成绩采集表导入导出（openpyxl）。

表格结构：考号 | 姓名 | 班级 | 学科1 | 学科2 ...
导入时按「考号」定位学生，空单元格表示缺考（写入 NULL）。
"""
from __future__ import annotations

import io
from typing import Optional

from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from ..models import Banji, Chengji, Kaohao, Kaoshi, KaoshiSet, Student, Subject


def export_template(kaoshi_id: int, banji_id: Optional[int], db) -> io.BytesIO:
    """导出Grade采集表 (含已有Grade)。"""
    kaoshi = db.get(Kaoshi, kaoshi_id)
    subjects = _subjects_of(kaoshi_id, db)

    q = select(Kaohao).where(Kaohao.kaoshi_id == kaoshi_id, Kaohao.delete_time.is_(None))
    if banji_id:
        q = q.where(Kaohao.banji_id == banji_id)
    kaohaos = db.scalars(q.order_by(Kaohao.school_id, Kaohao.banji_id, Kaohao.paixu)).all()

    cj_map = {}
    if kaohaos:
        rows = db.execute(
            select(Chengji.kaohao_id, Chengji.subject_id, Chengji.defen).where(
                Chengji.kaohao_id.in_([k.id for k in kaohaos])
            )
        ).all()
        for kh_id, sid, defen in rows:
            cj_map[(kh_id, sid)] = None if defen is None else float(defen)

    wb = Workbook()
    ws = wb.active
    ws.title = "Grade template"

    header = ["Exam no.", "Name", "Class"] + [s.title for s in subjects]
    ws.append(header)

    banji_cache: dict[int, str] = {}
    for kh in kaohaos:
        stu = db.get(Student, kh.student_id)
        if kh.banji_id not in banji_cache:
            bj = db.get(Banji, kh.banji_id)
            banji_cache[kh.banji_id] = bj.full_title if bj else str(kh.banji_id)
        row = [kh.kaohao or kh.id, stu.xingming if stu else "", banji_cache[kh.banji_id]]
        row += [cj_map.get((kh.id, s.id)) for s in subjects]
        ws.append(row)

    for i, w in enumerate([12, 10, 14] + [8] * len(subjects), start=1):
        ws.column_dimensions[chr(64 + i) if i <= 26 else "A"].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def import_scores(kaoshi_id: int, file_bytes: bytes, user_id: int, db) -> dict:
    """导入Grade。返回 {"success": str, "err": [str]}。"""
    wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {"success": "The sheet is empty", "err": []}

    header = [str(h).strip() if h else "" for h in rows[0]]
    subjects = _subjects_of(kaoshi_id, db)
    # 表头第 4 列起是学科名 -> subject_id
    col_subject: dict[int, int] = {}
    for idx, name in enumerate(header):
        if idx < 3:
            continue
        for s in subjects:
            if s.title == name:
                col_subject[idx] = s.id
                break

    if not col_subject:
        return {"success": "", "err": ["No matching subject column found in the header; please use the exported template."]}

    # 建立考号 -> Kaohao 的索引
    kh_index: dict[str, Kaohao] = {}
    for kh in db.scalars(
        select(Kaohao).where(Kaohao.kaoshi_id == kaoshi_id, Kaohao.delete_time.is_(None))
    ).all():
        kh_index[str(kh.kaohao or kh.id)] = kh
        kh_index[str(kh.id)] = kh

    ok_cnt = 0
    errs: list[str] = []
    for r in rows[1:]:
        if not r or r[0] is None:
            continue
        key = str(r[0]).strip()
        kh = kh_index.get(key)
        if kh is None:
            errs.append(f"Row {rows.index(r) + 1}: exam number {key} does not belong to this exam")
            continue

        for col, subject_id in col_subject.items():
            if col >= len(r):
                continue
            raw = r[col]
            defen: Optional[float] = None
            if raw is not None and str(raw).strip() != "":
                try:
                    defen = round(float(raw), 1)
                except (TypeError, ValueError):
                    errs.append(f"Exam number {key}: value {raw!r} in column {header[col]} is not a number")
                    continue
            _upsert(db, kh.id, subject_id, defen, user_id)
            ok_cnt += 1

    db.commit()
    return {"success": f"Imported {ok_cnt} grade rows", "err": errs}


def _upsert(db, kaohao_id: int, subject_id: int, defen: Optional[float], user_id: int):
    cj = db.scalars(
        select(Chengji).where(
            Chengji.kaohao_id == kaohao_id, Chengji.subject_id == subject_id
        )
    ).first()
    if cj is None:
        cj = Chengji(kaohao_id=kaohao_id, subject_id=subject_id, user_id=user_id)
        db.add(cj)
    cj.defen = defen
    cj.user_id = user_id
    # 得分率留待统计时统一计算，这里先清空避免脏数据
    cj.defenlv = None
    return cj


def _subjects_of(kaoshi_id: int, db) -> list[Subject]:
    """本次Exam涉及的Subject (按 kaoshi_set 去重, 保持Subject表Order)。"""
    sids = [
        r[0]
        for r in db.execute(
            select(KaoshiSet.subject_id)
            .where(KaoshiSet.kaoshi_id == kaoshi_id, KaoshiSet.delete_time.is_(None))
            .distinct()
        ).all()
    ]
    if not sids:
        return []
    return db.scalars(select(Subject).where(Subject.id.in_(sids)).order_by(Subject.paixu)).all()
