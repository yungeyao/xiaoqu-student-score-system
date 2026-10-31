"""Academic risk: 页面, 名单接口, 图表数据与导出。"""
from __future__ import annotations

import io
from collections import defaultdict

from fastapi import APIRouter, Depends, Request
from fastapi.responses import StreamingResponse
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..core.deps import get_db, require_admin
from ..core.response import fail, ok, table
from ..models import Banji, Kaoshi, Student
from ..services import risk as risk_service
from ..services.crud import int_or

router = APIRouter(tags=["risk"])


async def _form(request: Request) -> dict:
    return dict(await request.form())


def _exam_options(db):
    return [
        {"id": k.id, "title": k.title}
        for k in db.scalars(
            select(Kaoshi).where(Kaoshi.delete_time.is_(None)).order_by(Kaoshi.bfdate.desc())
        ).all()
    ]


def _banji_options(db, scope, keep_all: bool = True):
    stmt = select(Banji).where(Banji.delete_time.is_(None))
    if not scope["all"]:
        stmt = stmt.where(Banji.id.in_(scope["banji_ids"] or [0]))
    rows = db.scalars(stmt.order_by(Banji.school_id, Banji.ruxuenian, Banji.paixu)).all()
    return [{"id": b.id, "title": b.full_title} for b in rows]


def _scope(request, db):
    from ..services.permission import data_scope

    user, _cat, _r = require_admin(request, db)
    return data_scope(db, user)


def _banji_ids(scope) -> list[int] | None:
    """预警计算用的Class范围: 管理员为 None (不限)。"""
    return None if scope["all"] else (scope["banji_ids"] or [0])


# ================================================================== 页面


@router.get("/risk")
def risk_page(request: Request, db: Session = Depends(get_db)):
    user, cat, redirect = require_admin(request, db)
    if redirect:
        return redirect
    from ..main import templates

    scope = _scope(request, db)
    page = {
        "webtitle": "Academic risk",
        "exams": _exam_options(db),
        "banjis": _banji_options(db, scope),
    }
    return templates.TemplateResponse(request, "risk.html", {"request": request, "page": page})


# ================================================================== 数据


def _rows(request, db, form):
    kid = int_or(form, "kaoshi_id", 0)
    banji_id = int_or(form, "banji_id", 0)
    if not kid:
        return None, "Select an exam first"

    scope = _scope(request, db)
    ids = _banji_ids(scope)
    if banji_id:
        if ids is not None and banji_id not in ids:
            return None, "You cannot view this class"
        ids = [banji_id]

    rows = risk_service.evaluate(db, kid, ids)
    bj = {b.id: b.full_title for b in db.scalars(select(Banji)).all()}
    for r in rows:
        r["banji"] = bj.get(r["banji_id"], "")
    return rows, None


@router.post("/api/risk/list")
async def risk_list(request: Request, db: Session = Depends(get_db)):
    form = await _form(request)
    rows, err = _rows(request, db, form)
    if err:
        return fail(err)
    level = form.get("level") or ""
    if level:
        rows = [r for r in rows if r["risk_level"] == level]
    return table(len(rows), rows)


@router.post("/api/risk/summary")
async def risk_summary(request: Request, db: Session = Depends(get_db)):
    form = await _form(request)
    rows, err = _rows(request, db, form)
    if err:
        return fail(err)
    return ok("", risk_service.summary(rows))


@router.post("/api/risk/chart")
async def risk_chart(request: Request, db: Session = Depends(get_db)):
    """按Class分组的风险Level students数, 供柱状图使用。"""
    form = await _form(request)
    rows, err = _rows(request, db, form)
    if err:
        return fail(err)

    bj = {b.id: b.full_title for b in db.scalars(select(Banji)).all()}
    grouped: dict[int, dict[str, int]] = defaultdict(lambda: {lv: 0 for lv in risk_service.LEVELS})
    for r in rows:
        grouped[r["banji_id"]][r["risk_level"]] += 1

    order = sorted(grouped, key=lambda bid: (-grouped[bid]["High"], -grouped[bid]["Medium"]))
    categories = [bj.get(bid, "") for bid in order]
    series = [
        {"name": lv, "data": [grouped[bid][lv] for bid in order]}
        for lv in ("High", "Medium", "Watch")
    ]
    # 风险等级构成（饼图用）
    s = risk_service.summary(rows)
    pie = [{"name": lv, "value": s["counts"][lv]} for lv in risk_service.LEVELS]
    return ok("", {"categories": categories, "series": series, "pie": pie, "summary": s})


@router.get("/api/risk/export.xlsx")
def risk_export(request: Request, kaoshi_id: int = 0, banji_id: int = 0,
                db: Session = Depends(get_db)):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    user, cat, redirect = require_admin(request, db)
    if redirect:
        return redirect
    if not kaoshi_id:
        return fail("Select an exam first")

    form = {"kaoshi_id": str(kaoshi_id), "banji_id": str(banji_id or "")}
    rows, err = _rows(request, db, form)
    if err:
        return fail(err)
    summary = risk_service.summary(rows)

    wb = Workbook()
    ws = wb.active
    ws.title = "At-risk students"
    head = ["Class", "Name", "Total", "Average", "Class avg", "Gap", "Failed subjects",
            "Absent subjects", "Attendance (%)", "Risk score", "Level", "Reason"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1E9FFF")
        c.alignment = Alignment(horizontal="center")

    fills = {"High": "FFE7E7", "Medium": "FFF4E0", "Watch": "FFFBE0"}
    for r in rows:
        ws.append([
            r["banji"], r["xingming"], r["total"], r["avg"], r["class_avg"], r["gap"],
            r["fail_count"], r["absent_count"], r["attendance_rate"],
            r["risk_score"], r["risk_level"], "; ".join(r["reasons"]),
        ])
        color = fills.get(r["risk_level"])
        if color:
            for c in ws[ws.max_row]:
                c.fill = PatternFill("solid", fgColor=color)

    ws2 = wb.create_sheet("Summary")
    ws2.append(["Metric", "Value"])
    ws2.append(["Students analysed", summary["total"]])
    for lv in risk_service.LEVELS:
        ws2.append([lv, summary["counts"][lv]])
    ws2.append(["At-risk students", summary["at_risk"]])
    ws2.append(["At-risk share (%)", summary["at_risk_rate"]])

    for sheet in (ws, ws2):
        for column in sheet.columns:
            width = max(len(str(c.value or "")) for c in column) * 1.6 + 4
            sheet.column_dimensions[column[0].column_letter].width = min(width, 42)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    name = "academic-risk.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{name}"'},
    )
