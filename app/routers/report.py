"""报表中心：班级成绩报表、学生成绩单（可打印）与 Excel 批量导出。

报表不是把页面截图拼起来，而是同一次数据装配产出三种形态：
页面（打印/PDF）、单班工作簿、整场考试全部班级的批量工作簿。
"""
from __future__ import annotations

import io
from collections import defaultdict
from datetime import datetime

from fastapi import APIRouter, Depends, Request
from fastapi.responses import StreamingResponse
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..core.deps import get_db, require_admin
from ..core.response import fail
from ..models import (
    Banji,
    Chengji,
    Kaohao,
    Kaoshi,
    School,
    Student,
    Subject,
    TongjiBj,
    Xueqi,
)
from ..models.base import ts_to_str
from ..services import risk as risk_service
from ..services.permission import data_scope

router = APIRouter(tags=["report"])


def _exam(db, kid: int):
    return db.get(Kaoshi, kid) if kid else None


def _exam_options(db):
    return [
        {"id": k.id, "title": k.title}
        for k in db.scalars(
            select(Kaoshi).where(Kaoshi.delete_time.is_(None)).order_by(Kaoshi.bfdate.desc())
        ).all()
    ]


def _banji_options(db, scope):
    stmt = select(Banji).where(Banji.delete_time.is_(None))
    if not scope["all"]:
        stmt = stmt.where(Banji.id.in_(scope["banji_ids"] or [0]))
    rows = db.scalars(stmt.order_by(Banji.school_id, Banji.ruxuenian, Banji.paixu)).all()
    return [{"id": b.id, "title": b.full_title} for b in rows]


def _subjects(db):
    return db.scalars(
        select(Subject)
        .where(Subject.delete_time.is_(None), Subject.kaoshi == 1)
        .order_by(Subject.paixu, Subject.id)
    ).all()


def class_report(db, kid: int, banji_id: int) -> dict | None:
    """装配一个班的Grade report数据。"""
    exam = _exam(db, kid)
    banji = db.get(Banji, banji_id)
    if exam is None or banji is None or banji.delete_time is not None:
        return None

    subjects = _subjects(db)
    students = db.scalars(
        select(Student)
        .where(Student.banji_id == banji_id, Student.delete_time.is_(None))
        .order_by(Student.id)
    ).all()

    kh_rows = db.execute(
        select(Kaohao.id, Kaohao.student_id).where(
            Kaohao.kaoshi_id == kid, Kaohao.banji_id == banji_id
        )
    ).all()
    kh_to_student = {kh: sid for kh, sid in kh_rows}
    scores: dict[int, dict[int, dict]] = defaultdict(dict)
    if kh_to_student:
        rows = db.execute(
            select(
                Chengji.kaohao_id, Chengji.subject_id, Chengji.defen,
                Chengji.bpaixu, Chengji.biaozhunfen,
            ).where(Chengji.kaohao_id.in_(list(kh_to_student)))
        ).all()
        for kh_id, subject_id, defen, bpaixu, bzf in rows:
            sid = kh_to_student.get(kh_id)
            if sid is None:
                continue
            scores[sid][subject_id] = {
                "defen": None if defen is None else float(defen),
                "paixu": bpaixu,
                "biaozhunfen": None if bzf is None else float(bzf),
            }

    # 学生行 + 总分 + 班内排名
    rows_out = []
    for stu in students:
        cells = []
        total = 0.0
        complete = True
        for sub in subjects:
            cell = scores.get(stu.id, {}).get(sub.id)
            if cell is None or cell["defen"] is None:
                complete = False
                total = None
                cells.append({"defen": None, "paixu": None, "biaozhunfen": None})
                continue
            if total is not None:
                total = round(total + cell["defen"], 1)
            cells.append(cell)
        rows_out.append(
            {
                "student_id": stu.id,
                "xingming": stu.xingming,
                "shenfenzhenghao": stu.shenfenzhenghao,
                "banji_id": banji_id,
                "cells": cells,
                "total": total,
            }
        )

    ranked = sorted(
        [r for r in rows_out if r["total"] is not None],
        key=lambda r: r["total"],
        reverse=True,
    )
    rank_of: dict[int, int] = {}
    prev, prev_rank = None, 0
    for i, r in enumerate(ranked, start=1):
        if prev is None or r["total"] != prev:
            prev_rank = i
            prev = r["total"]
        rank_of[r["student_id"]] = prev_rank
    for r in rows_out:
        r["rank"] = rank_of.get(r["student_id"])

    # 报表按名次展示更直观，缺考学生排在最后
    rows_out.sort(key=lambda r: (r["rank"] is None, r["rank"] or 0))

    # 各科班级统计（直接取统计表，口径与统计页一致）
    stats = {}
    for t in db.scalars(
        select(TongjiBj).where(TongjiBj.kaoshi_id == kid, TongjiBj.banji_id == banji_id)
    ).all():
        stats[t.subject_id] = t

    # 预警
    risk_rows = risk_service.evaluate(db, kid, [banji_id], include_attendance=False)
    risk_by_student = {r["student_id"]: r for r in risk_rows}
    for r in rows_out:
        rr = risk_by_student.get(r["student_id"])
        r["risk_level"] = rr["risk_level"] if rr else "Normal"
        r["risk_reasons"] = "; ".join(rr["reasons"]) if rr else ""

    summary = risk_service.summary(risk_rows)
    total_avg = (
        round(sum(r["total"] for r in ranked) / len(ranked), 2) if ranked else None
    )
    school = db.get(School, banji.school_id) if banji.school_id else None

    return {
        "exam": exam,
        "banji": banji,
        "school_title": school.title if school else "",
        "students": rows_out,
        "subjects": subjects,
        "stats": stats,
        "subject_list": [stats.get(s.id) for s in subjects],
        "summary": summary,
        "total_avg": total_avg,
        "absent_total": sum(1 for r in rows_out if r["total"] is None),
        "exam_date": ts_to_str(exam.bfdate),
        "now_text": datetime.now().strftime("%Y-%m-%d %H:%M"),
    }


def student_report(db, kid: int, student_id: int) -> dict | None:
    exam = _exam(db, kid)
    stu = db.get(Student, student_id)
    if exam is None or stu is None or stu.delete_time is not None:
        return None
    banji = db.get(Banji, stu.banji_id)
    subjects = _subjects(db)

    kh = db.scalars(
        select(Kaohao).where(Kaohao.kaoshi_id == kid, Kaohao.student_id == student_id)
    ).first()

    cells = []
    total = 0.0
    for sub in subjects:
        row = None
        if kh is not None:
            row = db.scalars(
                select(Chengji).where(
                    Chengji.kaohao_id == kh.id, Chengji.subject_id == sub.id
                )
            ).first()
        defen = None if row is None or row.defen is None else float(row.defen)
        if defen is None:
            total = None
        elif total is not None:
            total = round(total + defen, 1)
        cells.append(
            {
                "subject": sub.title,
                "defen": defen,
                "paixu": row.bpaixu if row else None,
                "biaozhunfen": None if row is None or row.biaozhunfen is None
                else float(row.biaozhunfen),
            }
        )

    risk_rows = risk_service.evaluate(db, kid, [stu.banji_id], include_attendance=False)
    rr = next((r for r in risk_rows if r["student_id"] == student_id), None)

    xueqi = db.get(Xueqi, exam.xueqi_id) if exam.xueqi_id else None
    school = db.get(School, banji.school_id) if banji and banji.school_id else None
    return {
        "exam": exam,
        "exam_date": ts_to_str(exam.bfdate),
        "xueqi": xueqi.title if xueqi else "",
        "school": school.title if school else "",
        "banji": banji.full_title if banji else "",
        "student": stu,
        "kaohao": kh.kaohao if kh else "",
        "cells": cells,
        "total": total,
        "risk_level": rr["risk_level"] if rr else "Normal",
        "risk_reasons": "; ".join(rr["reasons"]) if rr else "",
    }


# ================================================================== 页面


@router.get("/report")
def report_page(request: Request, kaoshi_id: int = 0, db: Session = Depends(get_db)):
    user, cat, redirect = require_admin(request, db)
    if redirect:
        return redirect
    from ..main import templates

    scope = data_scope(db, user)
    page = {
        "webtitle": "Grade report",
        "exams": _exam_options(db),
        "banjis": _banji_options(db, scope),
        "preselect": kaoshi_id,
    }
    return templates.TemplateResponse(request, "report.html", {"request": request, "page": page})


@router.get("/report/banji")
def report_banji(request: Request, kaoshi_id: int = 0, banji_id: int = 0,
                 db: Session = Depends(get_db)):
    user, cat, redirect = require_admin(request, db)
    if redirect:
        return redirect
    scope = data_scope(db, user)
    if not scope["all"] and banji_id and banji_id not in scope["banji_ids"]:
        return fail("You cannot view this class")
    data = class_report(db, kaoshi_id, banji_id)
    if data is None:
        return fail("Exam or class not found")
    from ..main import templates

    return templates.TemplateResponse(
        request, "report_banji.html", {"request": request, "d": data}
    )


@router.get("/report/student")
def report_student(request: Request, kaoshi_id: int = 0, student_id: int = 0,
                   db: Session = Depends(get_db)):
    user, cat, redirect = require_admin(request, db)
    if redirect:
        return redirect
    data = student_report(db, kaoshi_id, student_id)
    if data is None:
        return fail("Exam or student not found")
    scope = data_scope(db, user)
    if not scope["all"] and data["student"].banji_id not in scope["banji_ids"]:
        return fail("You cannot view this student")
    from ..main import templates

    return templates.TemplateResponse(
        request, "report_student.html", {"request": request, "d": data}
    )


# ================================================================== Excel


def _style_sheet(ws, widths: dict[int, int] | None = None):
    from openpyxl.styles import Alignment, Font, PatternFill

    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1E9FFF")
        c.alignment = Alignment(horizontal="center")
    if widths:
        for idx, w in widths.items():
            ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = w


def _write_class_sheet(ws, data: dict):
    subjects = data["subjects"]
    head = ["Rank", "Name", "ID number"] + [s.title for s in subjects] + ["Total", "Risk"]
    ws.append(head)
    for r in data["students"]:
        ws.append(
            [r["rank"], r["xingming"], r["shenfenzhenghao"]]
            + [c["defen"] for c in r["cells"]]
            + [r["total"], r["risk_level"]]
        )
    _style_sheet(ws)


def _report_name(exam, banji=None, batch: bool = False) -> str:
    part = "batch" if batch else "class"
    return f"report-{part}-{exam.id}.xlsx"


@router.get("/api/report/banji.xlsx")
def report_banji_xlsx(request: Request, kaoshi_id: int = 0, banji_id: int = 0,
                      db: Session = Depends(get_db)):
    from openpyxl import Workbook

    user, cat, redirect = require_admin(request, db)
    if redirect:
        return redirect
    data = class_report(db, kaoshi_id, banji_id)
    if data is None:
        return fail("Exam or class not found")

    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    _write_class_sheet(ws, data)

    ws2 = wb.create_sheet("Subject statistics")
    ws2.append(["Subject", "Expected", "Taken", "Average", "Score rate", "Std Dev", "Max", "Min",
                "Excellent rate", "Pass rate", "Participation"])
    for sub, t in zip(data["subjects"], data["subject_list"]):
        if t is None:
            ws2.append([sub.title, None, None, None, None, None, None, None, None, None, None])
            continue
        ws2.append([
            sub.title, t.stu_cnt, t.chengji_cnt,
            float(t.avg) if t.avg is not None else None,
            float(t.defenlv) if t.defenlv is not None else None,
            float(t.biaozhuncha) if t.biaozhuncha is not None else None,
            float(t.max) if t.max is not None else None,
            float(t.min) if t.min is not None else None,
            float(t.youxiulv) if t.youxiulv is not None else None,
            float(t.jigelv) if t.jigelv is not None else None,
            float(t.canshilv) if t.canshilv is not None else None,
        ])
    _style_sheet(ws2)

    ws3 = wb.create_sheet("At-risk students")
    ws3.append(["Name", "Level", "Reason"])
    for r in data["students"]:
        if r["risk_level"] != "Normal":
            ws3.append([r["xingming"], r["risk_level"], r["risk_reasons"]])
    _style_sheet(ws3)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    name = _report_name(data["exam"])
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{name}"'},
    )


@router.get("/api/report/all.xlsx")
def report_all_xlsx(request: Request, kaoshi_id: int = 0, db: Session = Depends(get_db)):
    """一场Exam, 所有可见Class: 一个工作簿里每班一张表, 实现批量生成。"""
    from openpyxl import Workbook

    user, cat, redirect = require_admin(request, db)
    if redirect:
        return redirect
    exam = _exam(db, kaoshi_id)
    if exam is None:
        return fail("Exam not found")

    scope = data_scope(db, user)
    banjis = _banji_options(db, scope)
    if not banjis:
        return fail("No class to export")

    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Summary"
    ws0.append(["Class", "Taken", "Average total", "High", "Medium", "Watch"])
    total_generated = 0
    for b in banjis:
        data = class_report(db, kaoshi_id, b["id"])
        if data is None:
            continue
        ws0.append([
            b["title"], len(data["students"]) - data["absent_total"], data["total_avg"],
            data["summary"]["counts"]["High"], data["summary"]["counts"]["Medium"],
            data["summary"]["counts"]["Watch"],
        ])
        # 工作表名最长 31 字符，且不能含 : \ / ? * [ ]
        title = "".join(ch for ch in b["title"] if ch not in r':\/?*[]')[:28] or f"Class {b['id']}"
        ws = wb.create_sheet(title)
        _write_class_sheet(ws, data)
        total_generated += 1
    _style_sheet(ws0)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    name = f"report-batch-{exam.id}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{name}"',
            "X-Generated-Sheets": str(total_generated),
        },
    )
